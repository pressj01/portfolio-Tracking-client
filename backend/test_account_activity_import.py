"""Account-level broker activity recorded for whole-account performance."""

import datetime
import io
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import app as app_module
import config
import database


SCHWAB_CASH_ACTIVITY = "\n".join([
    "Date,Action,Symbol,Description,Quantity,Price,Fees & Comm,Amount",
    '04/01/2026,MoneyLink Transfer,,"Tfr BANK OF AMERICA, JANE DOE",,,,"$5,000.00"',
    "04/02/2026,Service Fee,,Account maintenance fee,,,,-$5.00",
    "04/03/2026,Wire Sent,,JANE DOE DISBURSED,,,,-$1000.00",
]).encode("utf-8")


class AccountActivityImportEndpointTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
        self.tmp.close()
        self.db_path = self.tmp.name
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        database.ensure_tables_exist(conn)
        conn.execute("DELETE FROM profiles")
        conn.execute(
            "INSERT INTO profiles (id, name, broker_source, include_in_owner, display_order) "
            "VALUES (1, 'Owner', '', 1, 1), (2, 'Schwab Brokerage', 'schwab', 1, 2)"
        )
        conn.commit()
        conn.close()

        self._orig_get_connection = app_module.get_connection
        self._orig_backup = app_module._create_import_backup
        self._orig_testing = app_module.app.testing
        self._orig_db_init = getattr(app_module.app, "_db_initialized", False)
        self._orig_db_path = config.DB_PATH
        config.DB_PATH = self.db_path
        app_module.get_connection = self._get_connection
        app_module._create_import_backup = lambda profile_id=None: None
        app_module.app.testing = True
        app_module.app._db_initialized = True
        self.client = app_module.app.test_client()

    def tearDown(self):
        app_module.get_connection = self._orig_get_connection
        app_module._create_import_backup = self._orig_backup
        app_module.app.testing = self._orig_testing
        app_module.app._db_initialized = self._orig_db_init
        config.DB_PATH = self._orig_db_path
        Path(self.db_path).unlink(missing_ok=True)

    def _get_connection(self):
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        return conn

    def _import(self):
        data = {
            "file": (io.BytesIO(SCHWAB_CASH_ACTIVITY), "Schwab_Transactions.csv"),
            "format": "schwab_transactions",
        }
        response = self.client.post(
            "/api/import/transactions?profile_id=2",
            data=data,
            content_type="multipart/form-data",
        )
        payload = response.get_json()
        self.assertEqual(response.status_code, 200, payload)
        return payload

    def test_import_records_each_signed_flow_once(self):
        first = self._import()
        second = self._import()

        self.assertEqual(first["account_activity_inserted"], 3)
        self.assertEqual(second["account_activity_inserted"], 0)
        self.assertEqual(second["account_activity_duplicates_skipped"], 3)
        self.assertIn("Nothing new to import", second["message"])

        conn = self._get_connection()
        try:
            rows = [
                tuple(row) for row in conn.execute(
                    "SELECT activity_date, activity_type, performance_treatment, base_amount "
                    "FROM account_activity WHERE profile_id = 2 ORDER BY activity_date"
                ).fetchall()
            ]
            owner_rows = conn.execute(
                "SELECT COUNT(*) FROM account_activity WHERE profile_id = 1"
            ).fetchone()[0]
        finally:
            conn.close()

        self.assertEqual(rows, [
            ("2026-04-01", "TRANSFER_IN", "EXTERNAL_FLOW", 5000.0),
            ("2026-04-02", "FEE", "EXPENSE", -5.0),
            ("2026-04-03", "WITHDRAWAL", "EXTERNAL_FLOW", -1000.0),
        ])
        self.assertEqual(owner_rows, 0)

    def test_only_files_that_carry_account_activity_record_coverage(self):
        # A dividends-only export is silent about deposits because it filtered
        # them out; its span must not read as "no deposits happened".
        dividends_only = "\n".join([
            "Date,Action,Symbol,Description,Quantity,Price,Fees & Comm,Amount",
            "03/02/2026,Cash Dividend,SCHD,SCHWAB US DIVIDEND EQUITY ETF,,,,$12.00",
        ]).encode("utf-8")
        response = self.client.post(
            "/api/import/transactions?profile_id=2",
            data={"file": (io.BytesIO(dividends_only), "Divs.csv"), "format": "schwab_transactions"},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        self._import()

        conn = self._get_connection()
        try:
            spans = [
                tuple(row) for row in conn.execute(
                    "SELECT start_date, end_date FROM account_activity_coverage WHERE profile_id = 2"
                ).fetchall()
            ]
        finally:
            conn.close()
        self.assertEqual(spans, [("2026-04-01", "2026-04-03")])

    def _call(self, method, url, body=None):
        response = getattr(self.client, method)(url, json=body)
        return response.status_code, response.get_json()

    def test_hand_entries_list_delete_and_guard(self):
        deposit = {"kind": "deposit", "date": "2026-04-01", "amount": "5,000", "note": "From checking"}
        self.assertEqual(self._call("post", "/api/account-activity?profile_id=2", deposit)[0], 200)
        # The same money on the same day again is refused unless confirmed.
        status, payload = self._call("post", "/api/account-activity?profile_id=2", deposit)
        self.assertEqual((status, payload.get("duplicate")), (409, True))
        status, _ = self._call("post", "/api/account-activity?profile_id=2", {**deposit, "allow_duplicate": True})
        self.assertEqual(status, 200)
        status, _ = self._call("post", "/api/account-activity?profile_id=2", {
            "kind": "transfer_in", "date": "2026-04-02", "ticker": "SCHD",
        })
        self.assertEqual(status, 400)
        status, added = self._call("post", "/api/account-activity?profile_id=2", {
            "kind": "transfer_out", "date": "2026-04-02", "ticker": "jepi", "quantity": 5,
        })
        self.assertEqual(status, 200)
        status, _ = self._call("post", "/api/account-activity/coverage?profile_id=2", {
            "start_date": "2026-04-01", "end_date": "2026-04-30",
        })
        self.assertEqual(status, 200)
        # Owner is a rollup here; writes must name the real account.
        status, _ = self._call("post", "/api/account-activity?profile_id=1", deposit)
        self.assertEqual(status, 400)

        status, listing = self._call("get", "/api/account-activity?profile_id=2")
        self.assertEqual(status, 200)
        self.assertTrue(listing["editable"])
        self.assertEqual(
            sorted((row["activity_type"], row["base_amount"], row["ticker"]) for row in listing["flows"]),
            [("DEPOSIT", 5000.0, None), ("DEPOSIT", 5000.0, None), ("SECURITY_TRANSFER_OUT", None, "JEPI")],
        )
        self.assertEqual(
            [(run["start_date"], run["end_date"]) for run in listing["covered_runs"]],
            [("2026-04-01", "2026-04-30")],
        )

        self.assertEqual(self._call("delete", f"/api/account-activity/{added['id']}?profile_id=2")[0], 200)
        self.assertEqual(self._call("delete", f"/api/account-activity/{added['id']}?profile_id=2")[0], 404)
        coverage_id = listing["coverage"][0]["id"]
        self.assertEqual(
            self._call("delete", f"/api/account-activity/coverage/{coverage_id}?profile_id=2")[0], 200,
        )

    def test_hand_entered_deposit_and_later_broker_import_stay_one_flow(self):
        self._call("post", "/api/account-activity?profile_id=2", {
            "kind": "deposit", "date": "2026-04-01", "amount": 5000,
        })

        payload = self._import()

        self.assertEqual(payload["account_activity_inserted"], 2)
        self.assertEqual(payload["account_activity_duplicates_skipped"], 1)

    def test_export_workbook_round_trips_activity_and_completed_periods(self):
        import openpyxl

        self._import()
        self._call("post", "/api/account-activity/coverage?profile_id=2", {
            "start_date": "2026-03-01", "end_date": "2026-03-31",
        })
        response = self.client.get("/api/export/holdings-transactions?profile_id=2")
        self.assertEqual(response.status_code, 200)
        workbook_bytes = response.data
        wb = openpyxl.load_workbook(io.BytesIO(workbook_bytes), read_only=True)
        try:
            self.assertIn("Account Activity", wb.sheetnames)
            self.assertIn("Activity Coverage", wb.sheetnames)
        finally:
            wb.close()

        conn = self._get_connection()
        try:
            conn.execute("DELETE FROM account_activity")
            conn.execute("DELETE FROM account_activity_coverage")
            conn.commit()
        finally:
            conn.close()

        restored = self.client.post(
            "/api/import/transactions?profile_id=2",
            data={
                "file": (io.BytesIO(workbook_bytes), "portfolio_with_transactions.xlsx"),
                "format": "portfolio_export",
                "export_scope": "transactions",
            },
            content_type="multipart/form-data",
        )
        payload = restored.get_json()
        self.assertEqual(restored.status_code, 200, payload)
        self.assertEqual(payload["account_activity_inserted"], 3)

        conn = self._get_connection()
        try:
            flows = [
                tuple(row) for row in conn.execute(
                    "SELECT activity_date, activity_type, base_amount FROM account_activity "
                    "WHERE profile_id = 2 ORDER BY activity_date"
                ).fetchall()
            ]
            spans = [
                tuple(row) for row in conn.execute(
                    "SELECT start_date, end_date FROM account_activity_coverage "
                    "WHERE profile_id = 2 ORDER BY start_date"
                ).fetchall()
            ]
        finally:
            conn.close()
        self.assertEqual(flows, [
            ("2026-04-01", "TRANSFER_IN", 5000.0),
            ("2026-04-02", "FEE", -5.0),
            ("2026-04-03", "WITHDRAWAL", -1000.0),
        ])
        self.assertEqual(spans, [("2026-03-01", "2026-03-31"), ("2026-04-01", "2026-04-03")])


class AccountAlphaEndpointTest(unittest.TestCase):
    """Account alpha from recorded values and imported flows, end to end."""

    END = datetime.date(2026, 8, 28)

    def setUp(self):
        self.tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
        self.tmp.close()
        self.db_path = self.tmp.name
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        database.ensure_tables_exist(conn)
        conn.execute("DELETE FROM profiles")
        conn.execute(
            "INSERT INTO profiles (id, name, broker_source, include_in_owner, display_order) "
            "VALUES (1, 'Owner', '', 0, 1), (2, 'Plain', 'schwab', 0, 2), "
            "(3, 'With Deposit', 'schwab', 0, 3)"
        )
        conn.commit()
        conn.close()

        days = [
            d.date() for d in pd.bdate_range(end=self.END, periods=120)
            if app_module.is_nyse_trading_day(d.date())
        ][-80:]
        rng = np.random.default_rng(7)
        spy_r = rng.normal(0.0004, 0.01, len(days) - 1)
        qqq_r = 0.6 * spy_r + rng.normal(0, 0.008, len(days) - 1)
        acct_r = 0.8 * spy_r + 0.0003 + rng.normal(0, 0.002, len(days) - 1)
        grow = lambda r: np.concatenate([[1.0], np.cumprod(1 + r)])  # noqa: E731
        index = pd.DatetimeIndex(days)
        self.days = days
        self.spy = pd.Series(400 * grow(spy_r), index=index)
        self.qqq = pd.Series(350 * grow(qqq_r), index=index)
        self.nav = pd.Series(100_000 * grow(acct_r), index=index)

        self.download_calls = 0

        def fake_download(tickers, **kwargs):
            self.download_calls += 1
            frame = pd.DataFrame({"SPY": self.spy, "QQQ": self.qqq})
            wanted = [t for t in tickers if t in frame.columns]
            return pd.concat({"Close": frame[wanted]}, axis=1)

        self._orig = (
            app_module.get_connection, app_module._chunked_yf_download,
            app_module.app.testing, getattr(app_module.app, "_db_initialized", False),
            config.DB_PATH,
        )
        config.DB_PATH = self.db_path
        app_module.get_connection = self._get_connection
        app_module._chunked_yf_download = fake_download
        app_module.app.testing = True
        app_module.app._db_initialized = True
        app_module._ACCOUNT_ALPHA_CACHE.clear()
        self.client = app_module.app.test_client()

    def tearDown(self):
        (
            app_module.get_connection, app_module._chunked_yf_download,
            app_module.app.testing, app_module.app._db_initialized, config.DB_PATH,
        ) = self._orig
        app_module._ACCOUNT_ALPHA_CACHE.clear()
        Path(self.db_path).unlink(missing_ok=True)

    def _get_connection(self):
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        return conn

    def _seed(self, profile_id, values, flows=(), covered=True):
        conn = self._get_connection()
        try:
            conn.executemany(
                "INSERT INTO portfolio_nav (profile_id, nav_date, total_value, source) VALUES (?, ?, ?, 'close')",
                [(profile_id, day.isoformat(), float(value)) for day, value in zip(self.days, values)],
            )
            for n, (day, amount) in enumerate(flows):
                conn.execute(
                    """INSERT INTO account_activity
                       (profile_id, activity_date, activity_type, direction, performance_treatment,
                        amount, base_amount, currency, dedupe_hash)
                       VALUES (?, ?, 'DEPOSIT', 'IN', 'EXTERNAL_FLOW', ?, ?, 'USD', ?)""",
                    (profile_id, day.isoformat(), amount, amount, f"test-{profile_id}-{n}"),
                )
            if covered:
                conn.execute(
                    "INSERT INTO account_activity_coverage (profile_id, start_date, end_date, source_format) "
                    "VALUES (?, ?, ?, 'schwab_transactions')",
                    (profile_id, self.days[0].isoformat(), self.days[-1].isoformat()),
                )
            conn.commit()
        finally:
            conn.close()

    def _alpha(self, profile_id):
        response = self.client.get(
            f"/api/account-alpha?profile_id={profile_id}&period=custom"
            f"&start_date={self.days[0].isoformat()}&end_date={self.days[-1].isoformat()}"
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        return response.get_json()

    def test_alpha_is_the_shared_regression_and_a_recorded_deposit_changes_nothing(self):
        self._seed(2, self.nav.values)
        # Same investment performance, plus $20k deposited on day 30 that then
        # grows with the account. Recorded as a flow, it must not read as return.
        deposit_day = self.days[30]
        growth = self.nav / self.nav.iloc[30]
        with_deposit = self.nav + (growth * 20_000).where(self.nav.index >= pd.Timestamp(deposit_day), 0)
        self._seed(3, with_deposit.values, flows=[(deposit_day, 20_000.0)])

        plain = self._alpha(2)
        deposited = self._alpha(3)

        expected = app_module._risk_profile(
            self.nav / self.nav.iloc[0], [("SPY", self.spy), ("QQQ", self.qqq)],
        )
        self.assertTrue(plain["available"], plain)
        self.assertEqual(plain["benchmark"], expected["beta_benchmark"])
        self.assertAlmostEqual(plain["alpha"], expected["alpha"], places=6)
        self.assertEqual(plain["observations"], len(self.days) - 1)
        self.assertAlmostEqual(deposited["alpha"], plain["alpha"], places=6)
        self.assertEqual(deposited["external_flow_count"], 1)

    def test_hidden_without_an_imported_activity_history(self):
        self._seed(2, self.nav.values, covered=False)

        payload = self._alpha(2)

        self.assertFalse(payload["available"])
        self.assertEqual(payload["reason"], "no_activity_history")
        self.assertEqual(self.download_calls, 0)

    def test_an_unrecorded_transfer_hides_the_number(self):
        jumped = self.nav.copy()
        jumped.iloc[40:] *= 1.6   # $60k arrives with no flow on file
        self._seed(2, jumped.values)

        payload = self._alpha(2)

        self.assertFalse(payload["available"])
        self.assertEqual(payload["reason"], "unexplained_jump")
        self.assertIn(self.days[40].isoformat(), payload["message"])


if __name__ == "__main__":
    unittest.main()
